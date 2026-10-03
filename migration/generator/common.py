"""أدوات مشتركة لمولّد ملفات النقل — تحويل القيم وكتابة SQL آمن."""
import datetime, re, json, warnings
from openpyxl import load_workbook

INVISIBLE = re.compile('[\u200B-\u200F\u202A-\u202E\u2066-\u2069\uFEFF\u00AD]')
AR_DIGITS = str.maketrans('٠١٢٣٤٥٦٧٨٩', '0123456789')


def norm_sheet(s):
    return ''.join(ch for ch in s if ch not in ' _').lower()


def load(path):
    with warnings.catch_warnings(record=True):
        warnings.simplefilter('always')
        wb = load_workbook(path, data_only=True)
    sheets = {}
    for ws in wb:
        rows, header = [], None
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            vals = list(row)
            if r_idx == 1:
                header = vals
                continue
            if all(v is None or (isinstance(v, str) and v == '') for v in vals):
                continue
            rows.append((r_idx, vals))
        sheets[norm_sheet(ws.title)] = {'title': ws.title, 'header': header, 'rows': rows}
    return sheets


def raw(v):
    """القيمة كنص حرفيًا كما في الشيت (بلا أي تنظيف) — للتخزين."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime.datetime):
        return v.isoformat()
    if isinstance(v, (datetime.date, datetime.time)):
        return v.isoformat()
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    return str(v)


def key(v):
    """نسخة منظّفة للمطابقة فقط (لا تُخزَّن): بلا مسافات أطراف ولا رموز خفية."""
    s = raw(v)
    if s is None:
        return ''
    return INVISIBLE.sub('', s).replace('\u00A0', ' ').strip()


def ar_norm(v):
    s = key(v)
    s = re.sub('[\u064B-\u065F\u0640]', '', s)
    s = re.sub('[أإآٱ]', 'ا', s).replace('ى', 'ي').replace('ة', 'ه')
    return re.sub(r'\s+', '', s).lower()


def split_list(v):
    return [x.strip() for x in key(v).split(',') if x.strip()]


def num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = key(v).translate(AR_DIGITS)
    return float(s) if re.fullmatch(r'-?\d+(\.\d+)?', s) else None


ORDINALS = [('اول', 1), ('ثاني', 2), ('ثالث', 3), ('رابع', 4), ('خامس', 5), ('سادس', 6), ('سابع', 7),
            ('ثامن', 8), ('تاسع', 9), ('عاشر', 10)]
TEENS = [('حادي', 11), ('ثاني', 12), ('ثالث', 13), ('رابع', 14), ('خامس', 15)]


def ordinal(v):
    """'الثاني' / 'الحادي عشر' / 2 / '٢' → رقم. None إن لم يُفهم."""
    n = num(v)
    if n is not None:
        return int(n) if n.is_integer() and n > 0 else None
    s = ar_norm(v)
    if not s:
        return None
    if 'عشر' in s:
        for w, k in TEENS:
            if w in s:
                return k
        return None
    for w, k in ORDINALS:
        if w in s:
            return k
    return None


def day_no(v):
    s = ar_norm(v)
    if not s:
        return None
    if 'حد' in s or s.startswith('sun'): return 0
    if 'ثنين' in s or 'تنين' in s or s.startswith('mon'): return 1
    if 'ثلاث' in s or 'تلات' in s or s.startswith('tue'): return 2
    if 'ربع' in s or s.startswith('wed'): return 3
    if 'خميس' in s or s.startswith('thu'): return 4
    return None


def period_no(v):
    n = ordinal(v)
    return n if n is not None and 1 <= n <= 12 else None


def parse_time(v):
    """datetime.time أو '10:30' أو '١٠:٣٠ ص' → 'HH:MM'"""
    if v is None or v == '':
        return None
    if isinstance(v, datetime.time):
        return v.strftime('%H:%M')
    if isinstance(v, datetime.datetime):
        return v.strftime('%H:%M')
    s = key(v).translate(AR_DIGITS).lower()
    m = re.match(r'(\d{1,2}):(\d{2})', s)
    if not m:
        return None
    h, mi = int(m.group(1)), int(m.group(2))
    pm = 'م' in s or 'pm' in s
    am = 'ص' in s or 'am' in s
    if pm and h < 12: h += 12
    if am and h == 12: h = 0
    return f'{h:02d}:{mi:02d}' if h < 24 and mi < 60 else None


def parse_ts(v):
    """قيمة تاريخ/وقت → نص timestamptz. التواريخ بلا منطقة تُعامل بتوقيت الرياض."""
    if v is None or v == '':
        return None
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d %H:%M:%S') + '+03'
    if isinstance(v, datetime.date):
        return v.isoformat() + ' 00:00:00+03'
    s = key(v)
    if re.match(r'^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}', s):
        return s  # ISO كامل (غالبًا بـ Z) — PostgreSQL يفهمه كما هو
    m = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T](\d{1,2}):(\d{2}))?', s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        h, mi = int(m.group(4) or 0), int(m.group(5) or 0)
        return f'{y:04d}-{mo:02d}-{d:02d} {h:02d}:{mi:02d}:00+03'
    return None


def parse_date(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    s = key(v)
    m = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})', s)
    if m and 1900 < int(m.group(1)) < 2100:
        return f'{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}'
    return None


def lit(v):
    """قيمة → SQL literal آمن."""
    if v is None:
        return 'null'
    if isinstance(v, bool):
        return 'true' if v else 'false'
    if isinstance(v, (int, float)):
        return repr(int(v)) if float(v).is_integer() else repr(v)
    s = str(v)
    # ⚠️ نصوص فيها أسطر جديدة أو رموز تحكم تُكتب بصيغة E'...' مع رموز هروب (\n)، لا كأسطر حقيقية:
    # النسخ واللصق في المتصفح على Windows يحوّل السطر الحقيقي داخل النص إلى \r\n فيتغير المحتوى.
    if any(ch in s for ch in '\n\r\t') or any(ord(ch) < 32 for ch in s):
        esc = s.replace('\\', '\\\\').replace("'", "\\'").replace('\n', '\\n').replace('\r', '\\r').replace('\t', '\\t')
        esc = ''.join(ch if ord(ch) >= 32 else '\\x%02x' % ord(ch) for ch in esc)
        return "E'" + esc + "'"
    return "'" + s.replace("'", "''") + "'"


def jlit(obj):
    return lit(json.dumps(obj, ensure_ascii=False)) + '::jsonb'


class Batch:
    """ملف SQL واحد = معاملة واحدة، مع محاسبة كل صف وتحقق ذاتي في النهاية."""

    def __init__(self, name, label):
        self.name, self.label = name, label
        self.lines = []
        self.expect = {}      # جدول → عدد متوقع
        self.accounted = {}   # شيت → عدد الصفوف المحاسَبة في هذا الملف
        self.quarantine = []

    def sql(self, s):
        self.lines.append(s)

    def comment(self, s):
        self.lines.append('\n-- ' + '-' * 66 + '\n-- ' + s + '\n-- ' + '-' * 66)

    def count(self, table, n=1):
        self.expect[table] = self.expect.get(table, 0) + n

    def account(self, sheet, row, status, target_table=None, target_id=None, note=None):
        self.accounted[sheet] = self.accounted.get(sheet, 0) + 1
        self.lines.append(
            'insert into migration.row_map (source_sheet, source_row, status, target_table, target_id, note) values ('
            + ', '.join(lit(x) for x in [sheet, row, status, target_table, target_id, note]) + ');')

    def quarantine_row(self, sheet, row, column, value, reason):
        self.quarantine.append((sheet, row, column, value, reason))
        self.lines.append(
            "insert into migration.quarantine (run_id, source_sheet, source_row, source_column, raw_value, reason) values ("
            "(select max(id) from migration.runs), " + ', '.join(lit(x) for x in [sheet, row, column, value, reason]) + ');')

    def render(self, baseline_counts):
        out = [f'-- =====================================================================',
               f'-- مِرقاة — نقل البيانات: {self.label}',
               f'-- معاملة واحدة: إما يُكتب كل شيء أو لا شيء. التحقق في النهاية يُلغي كل شيء عند أي اختلاف.',
               f'-- لا تشغّل هذا الملف مرتين (سيرفضه تكرار البيانات تلقائيًا).',
               f'-- =====================================================================',
               'begin;',
               "set local app.migrating = 'on';",
               f"insert into migration.runs (source_label, notes) values ({lit(self.name)}, {lit(self.label)});"]
        out += self.lines
        out.append('\n-- ' + '=' * 66 + '\n-- التحقق الذاتي: أي اختلاف يرمي خطأ فتُلغى المعاملة كاملة\n-- ' + '=' * 66)
        checks = []
        for table, n in sorted(self.expect.items()):
            base = baseline_counts.get(table, 0)
            checks.append(f"  if (select count(*) from {table}) <> {base + n} then raise exception "
                          f"'تحقق فشل: {table} المتوقع {base + n} الموجود %', (select count(*) from {table}); end if;")
        for sheet, n in sorted(self.accounted.items()):
            checks.append(f"  if (select count(*) from migration.row_map where source_sheet = {lit(sheet)}) <> {baseline_counts.get('rowmap:' + sheet, 0) + n} "
                          f"then raise exception 'تحقق فشل: محاسبة صفوف {sheet}'; end if;")
        out.append('do $$ begin\n' + '\n'.join(checks) + "\n  raise notice '✅ تحقق ناجح: " + self.label + "';\nend $$;")
        out.append("update migration.runs set finished_at = now(), status = 'succeeded' where id = (select max(id) from migration.runs);")
        out.append('commit;')
        return '\n'.join(out) + '\n'
